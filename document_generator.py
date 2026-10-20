"""Render resume text as a formatted, two-column Microsoft Word document."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.shared import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_color):
    """Set a table cell's background fill using WordprocessingML.

    Args:
        cell (docx.table._Cell): The table cell to format.
        fill_color (str): The fill color as a hexadecimal RGB string without
            ``#``.

    Returns:
        None: This function does not return a value.
    """
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def _add_formatted_text(paragraph, text, font_color=None):
    """Add text runs with Markdown-style bold markers and an optional color.

    Text between successive pairs of ``**`` markers is rendered in bold; the
    marker characters themselves are omitted.

    Args:
        paragraph (docx.text.paragraph.Paragraph): The paragraph that receives
            the formatted runs.
        text (str): Plain text with optional ``**bold**`` segments.
        font_color (RGBColor | None): Optional color applied to every
            generated run.

    Returns:
        None: This function does not return a value.
    """
    parts = text.split('**')
    for i, part in enumerate(parts):
        run = paragraph.add_run(part)
        if font_color:
            run.font.color.rgb = font_color
        if i % 2 != 0:
            run.bold = True

def save_as_word(text, filename="generated_resume.docx"):
    """Create a styled two-column resume document from structured text.

    The left and right columns are separated by ``[JOBB_OSZLOP]``. If that
    marker is absent, all supplied text is placed in the left column.

    Args:
        text (str): Resume content, optionally containing the column marker
            and Markdown-like headings, lists, and bold spans.
        filename (str): Destination path for the generated ``.docx`` file.

    Returns:
        None: This function does not return a value.
    """
    doc = Document()
    
    # 1. Set compact margins for a modern layout.
    for section in doc.sections:
        section.top_margin = Cm(1.0)
        section.bottom_margin = Cm(1.0)
        section.left_margin = Cm(1.0)
        section.right_margin = Cm(1.0)
        
    # 2. Split the content into the two resume columns.
    left_text = ""
    right_text = ""
    if "[JOBB_OSZLOP]" in text:
        parts = text.split("[JOBB_OSZLOP]")
        left_text = parts[0].replace("[BAL_OSZLOP]", "").strip()
        right_text = parts[1].strip()
    else:
        left_text = text # Fallback when no column separator is provided
        
    # 3. Build the two-column layout with a table.
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    
    cell_left = table.cell(0, 0)
    cell_right = table.cell(0, 1)
    
    # Set column widths to approximately 35% left and 65% right.
    cell_left.width = Cm(6.5)
    cell_right.width = Cm(12.5)
    
    # Apply a dark gray background to the left column.
    set_cell_background(cell_left, '333333')
    
    def write_to_cell(md_text, cell, is_left=False):
        """Render resume lines into a table cell with column-specific styling.

        Args:
            md_text (str): Text content to render in the cell.
            cell (docx.table._Cell): Destination table cell.
            is_left (bool): Whether to apply the left-column layout and light
                text.

        Returns:
            None: This function does not return a value.
        """
        # Use light text on the dark left column and dark text on the right.
        text_color = RGBColor(255, 255, 255) if is_left else RGBColor(0, 0, 0)
        
        if is_left:
            # Add a photo placeholder at the top of the left column.
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run("\n[ 📷 FÉNYKÉP ]\n(Kattints ide)\n\n")
            run.font.color.rgb = text_color
            run.bold = True
            
        lines = md_text.split('\n')
        for line in lines:
            line = line.strip()
            if not line:
                continue
                
            if line.startswith('# '): # Name
                p = cell.add_paragraph()
                r = p.add_run(line.replace('# ', '').upper())
                r.bold = True
                r.font.size = Pt(16)
                r.font.color.rgb = text_color
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(12)
            elif line.startswith('## '): # Section headings
                p = cell.add_paragraph()
                r = p.add_run(line.replace('## ', '').upper())
                r.bold = True
                r.font.size = Pt(12)
                r.font.color.rgb = text_color
                p.paragraph_format.space_before = Pt(10)
            elif line.startswith('- ') or line.startswith('* '): # Lists
                p = cell.add_paragraph()
                _add_formatted_text(p, "• " + line[2:], font_color=text_color)
                p.paragraph_format.left_indent = Cm(0.5) # Indent list items.
            else: # Plain text, such as a professional summary.
                p = cell.add_paragraph()
                _add_formatted_text(p, line, font_color=text_color)
                
    # Render the content into each column.
    write_to_cell(left_text, cell_left, is_left=True)
    write_to_cell(right_text, cell_right, is_left=False)
    
    doc.save(filename)