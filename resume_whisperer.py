import os
import time
from openai import OpenAI
from dotenv import load_dotenv
from docx import Document

load_dotenv()
API_KEY = os.getenv("OPENAI_API_KEY")

if not API_KEY:
    print("❌ Hiba: Nem találom az OPENAI_API_KEY-t a .env fájlban!")
    exit()

client = OpenAI(api_key=API_KEY)

system_prompt = """
You are a professional, multilingual HR assistant and career coach.
Rules of the interview:
1. The user will state their PREFERRED LANGUAGE for the final resume. Remember this!
2. Your goal is to gather detailed information about Education, Experience, and Skills. Ask in HUNGARIAN.
3. OVERRIDE RULE: If the user provides a comprehensive summary of their Education, Experience, and Skills in a single prompt, DO NOT ask any follow-up questions! Accept it as complete and IMMEDIATELY proceed to generation.
4. DRILL-DOWN RULE: Only ask follow-up questions if the user's answer is extremely short or completely lacks context (e.g., "I worked at a shop"). If they provide a detailed paragraph, accept it without asking for more!
5. Once you have the data, STOP asking questions.
6. Generate the final resume strictly in the PREFERRED TARGET LANGUAGE requested by the user.
7. IMPORTANT FORMATTING: At the very top of the generated resume, you MUST include this exact placeholder for the photo: "[ 📷 FÉNYKÉP HELYE ]"
8. You MUST start your final message with the exact word 'RESUME_GENERATED'.
9. STRICT RULE: After 'RESUME_GENERATED', you MUST output ONLY the raw Markdown text of the resume. DO NOT include any conversational filler, greetings, explanations, or closing remarks like "Here is your resume". JUST the pure resume content.
"""
kezdoszoveg = "Szia! Egy többnyelvű AI karrier-coach vagyok. Első lépésként kérlek mondd el: milyen nyelven szeretnéd megkapni a kész önéletrajzodat a beszélgetés végén?"
messages = [
    {"role": "system", "content": system_prompt},
    {"role": "assistant", "content": kezdoszoveg}
]

def record_audio_to_text():
    return input("\n⌨️ Te (gépelj ide): ")

def _add_formatted_text(paragraph, text):
    """Segédfüggvény a **félkövér** szövegek szép megjelenítéséhez Wordben"""
    parts = text.split('**')
    for i, part in enumerate(parts):
        run = paragraph.add_run(part)
        if i % 2 != 0:  # A páratlan indexű elemek a ** között voltak a Markdownban
            run.bold = True

def save_as_word(text, filename="generated_resume.docx"):
    """Továbbfejlesztett Word generáló, ami sokkal szebben formáz"""
    doc = Document()
    
    lines = text.split('\n')
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        if line.startswith('# '):
            doc.add_heading(line.replace('# ', ''), level=1)
        elif line.startswith('## '):
            doc.add_heading(line.replace('## ', ''), level=2)
        elif line.startswith('### '):
            doc.add_heading(line.replace('### ', ''), level=3)
        elif line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            _add_formatted_text(p, line[2:])
        else:
            p = doc.add_paragraph()
            _add_formatted_text(p, line)
            
    doc.save(filename)

def chat_with_ai(user_text):
    messages.append({"role": "user", "content": user_text})
    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=messages
    )
    ai_reply = response.choices[0].message.content
    messages.append({"role": "assistant", "content": ai_reply})
    return ai_reply

print("=== Üdvözöl a Resume Whisperer! (SZÖVEGES TESZT MÓD) ===")
print("Indítom a HR interjút...\n")
print(f"🤖 AI: {kezdoszoveg}")

while True:
    user_input = None
    while not user_input:
        user_input = record_audio_to_text()
    
    ai_response = chat_with_ai(user_input)
    
    if "RESUME_GENERATED" in ai_response:
        clean_resume = ai_response.replace("RESUME_GENERATED", "").strip()
        
        # Word generálása az új, formázott logikával
        save_as_word(clean_resume, "generated_resume.docx")
        
        with open("generated_resume.md", "w", encoding="utf-8") as f:
            f.write(clean_resume)
            
        print("\n🤖 AI: Szuper, a cv-d le lett generálva, már csak egy fényképet kell beszúrnod, és kész is vagy!")
        break
    else:
        print(f"\n🤖 AI: {ai_response}")
    
    time.sleep(1)