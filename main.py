import time
import json
from docx import Document
from audio_handler import record_audio_to_text
from ai_handler import chat_with_ai, kezdoszoveg, init_ai
from document_generator import generate_word_with_semantic_mapping

print("=== Üdvözöl a Resume Whisperer! ===")

# ÚJ LÉPÉS: Gyors sablon olvasás
print("🔄 Sablon kapacitásának felmérése...")
doc = Document("CVtemplate.docx")
template_text_lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            template_text_lines.extend([p.text.strip() for p in cell.paragraphs if p.text.strip()])

# Odaaduk a nyers szöveget az AI-nak, hogy képbe kerüljön a dizájnnal
init_ai("\n".join(template_text_lines))

print("✅ Sablon betöltve! Indítom a HR interjút...\n")
print(f"🤖 AI: {kezdoszoveg}")

while True:
    user_input = None
    while not user_input:
        user_input = record_audio_to_text()
    
    ai_response_json_string = chat_with_ai(user_input)
    
    try:
        ai_data = json.loads(ai_response_json_string)
        
        if ai_data.get("status") == "RESUME_GENERATED":
            print("\n🤖 AI: Szuper, összegyűjtöttem mindent! Töltöm a sablonba...")
            
            # Intelligens AI Sablon-töltő meghívása
            generate_word_with_semantic_mapping(ai_data, "CVtemplate.docx", "generated_resume.docx")
            
            # Elmentjük a json-t a demóhoz
            with open("resume_data.json", "w", encoding="utf-8") as f:
                json.dump(ai_data, f, ensure_ascii=False, indent=4)
                
            print("✅ A cv-d le lett generálva a hozott formátumban! Nézd meg a generated_resume.docx fájlt!")
            break
        else:
            print(f"\n🤖 AI: {ai_data.get('message', '...')}")
            
    except json.JSONDecodeError:
        print("\n❌ Hiba: Az AI nem érvényes JSON-t küldött vissza. Próbáljuk újra!")
    
    time.sleep(1)